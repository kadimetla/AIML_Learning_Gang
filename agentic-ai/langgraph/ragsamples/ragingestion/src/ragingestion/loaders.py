"""Step 1 - LOAD: turn a file into LangChain Documents via a pluggable loader registry.

Each loader is registered under a name and the extensions it handles. The first loader
registered for an extension is its default; pass `loader="name"` to pick another one.
Add your own with `@register("name", {".ext"})` - heavy imports belong inside the function
so optional dependencies are only needed when that loader is actually used.
"""

from collections.abc import Callable
from pathlib import Path

from langchain_core.documents import Document

Loader = Callable[[Path, str], list[Document]]

_LOADERS: dict[str, tuple[Loader, frozenset[str]]] = {}
_DEFAULTS: dict[str, str] = {}  # extension -> default loader name


def register(name: str, extensions: set[str]):
    def deco(fn: Loader) -> Loader:
        _LOADERS[name] = (fn, frozenset(extensions))
        for ext in extensions:
            _DEFAULTS.setdefault(ext, name)
        return fn

    return deco


def available_loaders() -> dict[str, list[str]]:
    return {name: sorted(exts) for name, (_, exts) in _LOADERS.items()}


def supported_extensions() -> list[str]:
    return sorted(_DEFAULTS)


TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".py", ".html", ".htm"}


@register("pypdf", {".pdf"})
def _pypdf(path: Path, source: str) -> list[Document]:
    from pypdf import PdfReader

    docs = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": source, "page": i}))
    return docs


@register("pymupdf4llm", {".pdf"})
def _pymupdf4llm(path: Path, source: str) -> list[Document]:
    """Layout-aware PDF -> Markdown (keeps headings/tables); pair with the `markdown` chunk strategy."""
    import pymupdf4llm

    pages = pymupdf4llm.to_markdown(str(path), page_chunks=True)
    return [
        Document(page_content=p["text"], metadata={"source": source, "page": p["metadata"]["page"] + 1})
        for p in pages
        if p["text"].strip()
    ]


@register("text", TEXT_EXT)
def _text(path: Path, source: str) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": source})]


@register("docx", {".docx"})
def _docx(path: Path, source: str) -> list[Document]:
    import docx

    text = "\n\n".join(p.text for p in docx.Document(str(path)).paragraphs if p.text.strip())
    return [Document(page_content=text, metadata={"source": source})]


def load_file(path: Path, source_name: str | None = None, loader: str | None = None) -> list[Document]:
    source = source_name or path.name
    ext = path.suffix.lower()
    if loader is None:
        loader = _DEFAULTS.get(ext)
        if loader is None:
            raise ValueError(f"Unsupported file type {ext!r}. Supported: {' '.join(supported_extensions())}")
    elif loader not in _LOADERS:
        raise ValueError(f"Unknown loader {loader!r}; choose from {list(_LOADERS)}")
    fn, exts = _LOADERS[loader]
    if ext not in exts:
        raise ValueError(f"Loader {loader!r} does not handle {ext!r} (handles {sorted(exts)})")
    return fn(path, source)
